#!/usr/bin/env python3
# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 Dennis Rudin - see LICENSE and NOTICE
"""coursekit - turn a slide deck and a narration script into an audiobook and a video course.

    python coursekit.py setup [en] [sv] [piper:VOICE]   first run: download voice models (default: en, the Kokoro voices)
    python coursekit.py skill [--code]                  build the Claude skill: dist/slide-video-course.zip
                                                        (--code: also install it for Claude Code in ~/.claude/skills)
    python coursekit.py init  deck.pptx [--course DIR] [--lang en|sv]   slides + notes -> a new course folder
    python coursekit.py check DIR                       check the scripts against the slides
    python coursekit.py voicetest DIR                   speak a test sentence -> DIR/out/voice_test.wav
    python coursekit.py audio DIR [01 02 ...]           narration -> out/audio/*.mp3, *.m4b, *.timing.json
    python coursekit.py video DIR [01 02 ...]           slides + narration -> out/video/*.mp4 (+ YouTube chapters)
    python coursekit.py subs  DIR                       subtitles -> out/subtitles/srt/*.LANG.srt
    python coursekit.py all   DIR                       audio, video and subtitles
Add --budget SECONDS to audio, video or all to stop after that long; run the same command again to continue.
Everything here runs locally without AI. Writing the script and translating subtitles is done by Claude
(see the coursekit skill); the formats are described in README.md.
"""
import argparse, glob, hashlib, json, os, re, shutil, subprocess, sys, tempfile, time, wave

KIT = os.path.dirname(os.path.abspath(__file__))
MODELS = os.path.join(KIT, 'models')
SR = 44100
GAP_PARA, GAP_HEAD, GAP_CLIP = 0.55, 1.0, 1.0
SPEECH_LUFS = -16.0
FPS = 5
FOREIGN_GAP = 0.25
DEFAULTS = {
    'en': dict(engine='kokoro', voice='bm_george', speed=0.95),
    'sv': dict(engine='piper', voice='sv_SE-nst-medium', speed=1.0),
}
KOKORO_LANG = {'en': 'en-gb', 'en-us': 'en-us', 'sv': 'sv', 'fr': 'fr-fr', 'de': 'de', 'it': 'it', 'es': 'es',
               'pt': 'pt-br', 'cs': 'cs', 'no': 'nb', 'da': 'da', 'fi': 'fi', 'nl': 'nl', 'la': 'la'}
TRACK_WORD = {'en': 'audio', 'sv': 'ljud', 'de': 'Audio', 'es': 'audio', 'fr': 'audio'}
started = time.time()

class Paused(Exception): pass

def over_budget(a):
    return a.budget and time.time() - started > a.budget

# ---------------------------------------------------------------- course files
def load_course(d):
    p = os.path.join(d, 'course.json')
    if not os.path.exists(p): sys.exit(f'{p} not found - run "coursekit.py init" first')
    c = json.load(open(p, encoding='utf-8')); c['dir'] = os.path.abspath(d); return c

def modules(c, only=()):
    ms = sorted(glob.glob(os.path.join(c['dir'], 'script', '*.md')))
    if not ms: sys.exit(f'no scripts in {c["dir"]}/script/ yet')
    return [m for m in ms if not only or os.path.basename(m).split('_')[0] in only or os.path.basename(m)[:-3] in only]

def out(c, *p):
    f = os.path.join(c['dir'], 'out', *p); os.makedirs(os.path.dirname(f), exist_ok=True); return f

def parse(path):
    """Script -> blocks: ('head', (level, text)) ('say', text) ('pause', s) ('slide', n) ('clip', file)."""
    text = re.sub(r'<!--.*?-->', '', open(path, encoding='utf-8').read(), flags=re.S)
    blocks, para = [], []
    def flush():
        if para: blocks.append(('say', ' '.join(para))); para.clear()
    for line in text.splitlines():
        l = line.strip()
        m = re.fullmatch(r'\[(SLIDE|PAUSE|CLIP)\s+([^\]]+)\]', l)
        if not l: flush()
        elif m:
            flush(); k, v = m[1].lower(), m[2].strip()
            blocks.append((k, int(v) if k == 'slide' else float(v) if k == 'pause' else v))
        elif l.startswith('#'): flush(); blocks.append(('head', (len(l) - len(l.lstrip('#')), l.lstrip('#').strip())))
        else: para.append(l)
    flush(); return blocks

plain = lambda s: re.sub(r'\{[\w-]+\|([^{}]*)\}', r'\1', s)

def run(*cmd, **kw):
    return subprocess.run([str(x) for x in cmd], check=True, **kw)

def decode(path):
    import numpy as np
    raw = subprocess.run(['ffmpeg', '-loglevel', 'error', '-i', path, '-f', 'f32le', '-ac', '2', '-ar', str(SR), '-'],
                         capture_output=True, check=True).stdout
    return np.frombuffer(raw, dtype=np.float32).reshape(-1, 2)

def lufs_of(path):
    r = subprocess.run(['ffmpeg', '-nostats', '-i', path, '-af', 'ebur128', '-f', 'null', '-'], capture_output=True, text=True)
    return float(re.findall(r'I:\s+(-?[\d.]+) LUFS', r.stderr)[-1])

def download(url, dest):
    import urllib.request
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    print(f'downloading {os.path.basename(dest)} ...', flush=True)
    with urllib.request.urlopen(url, timeout=600) as r, open(dest + '.part', 'wb') as f: shutil.copyfileobj(r, f)
    os.replace(dest + '.part', dest)

# ---------------------------------------------------------------- init: slides and notes
def find_soffice():
    for s in ('soffice', 'libreoffice', '/Applications/LibreOffice.app/Contents/MacOS/soffice'):
        if shutil.which(s) or os.path.exists(s): return shutil.which(s) or s

def cmd_init(a):
    deck = os.path.abspath(a.deck); stem, ext = os.path.splitext(deck); ext = ext.lower()
    d = os.path.abspath(a.course or stem + '_course'); os.makedirs(os.path.join(d, 'slides'), exist_ok=True)
    for sub in ('script', 'source', 'media'): os.makedirs(os.path.join(d, sub), exist_ok=True)
    pdf = stem + '.pdf' if ext == '.pptx' and os.path.exists(stem + '.pdf') else deck if ext == '.pdf' else None
    if ext == '.pptx' and not pdf:
        so = find_soffice()
        if not so: sys.exit('To draw the slides I need a PDF of the deck: in PowerPoint choose File > Export > PDF and save it '
                            f'next to the deck as {os.path.basename(stem)}.pdf (or install LibreOffice), then run init again.')
        tmp = tempfile.mkdtemp(); run(so, '--headless', '--convert-to', 'pdf', '--outdir', tmp, deck, capture_output=True)
        pdf = os.path.join(tmp, os.path.basename(stem) + '.pdf')
    print('drawing slides ...', flush=True)
    run('pdftoppm', '-png', '-scale-to', 1920, pdf, os.path.join(d, 'slides', 's'))
    for f in glob.glob(os.path.join(d, 'slides', 's-*.png')):         # s-1.png / s-01.png -> s-001.png
        n = int(re.search(r's-(\d+)\.png$', f)[1]); os.replace(f, os.path.join(d, 'slides', f's-{n:03d}.png'))
    n_pages = len(glob.glob(os.path.join(d, 'slides', 's-*.png')))
    slides = []
    if ext == '.pptx':
        from pptx import Presentation
        for i, s in enumerate(Presentation(deck).slides, 1):
            title = s.shapes.title.text_frame.text.strip() if s.shapes.title is not None and s.shapes.title.has_text_frame else ''
            texts = [sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide and s.notes_slide.notes_text_frame else ''
            slides.append(dict(slide=i, title=title, text='\n'.join(texts), notes=notes, hidden=s._element.get('show') == '0'))
    else:
        for i in range(1, n_pages + 1):
            t = subprocess.run(['pdftotext', '-layout', '-f', str(i), '-l', str(i), pdf, '-'], capture_output=True, text=True).stdout
            slides.append(dict(slide=i, title=t.strip().split('\n')[0].strip() if t.strip() else '', text=t.strip(), notes='', hidden=False))
    json.dump(slides, open(os.path.join(d, 'source', 'slides.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    with open(os.path.join(d, 'source', 'slides.md'), 'w', encoding='utf-8') as f:   # easy to read for the script writer
        for s in slides:
            f.write(f"## Slide {s['slide']}: {s['title']}{'  (hidden)' if s['hidden'] else ''}\n\nOn the slide:\n{s['text']}\n\n"
                    f"Teacher's notes:\n{s['notes'] or '(none)'}\n\n")
    cfg = os.path.join(d, 'course.json')
    if not os.path.exists(cfg):
        lang = a.lang; c = dict(title=os.path.basename(stem).replace('_', ' '), author='', language=lang, deck=deck,
                                **DEFAULTS.get(lang, DEFAULTS['en']), subtitles=[lang],
                                voicebox_url='http://127.0.0.1:17493', voicebox_profile='')
        json.dump(c, open(cfg, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'{len(slides)} slides ({n_pages} drawn) -> {d}\n  slides/s-NNN.png, source/slides.md (text + notes), course.json\n'
          f'Next: write the narration in {d}/script/ (see README.md), then "coursekit.py check {d}".')

# ---------------------------------------------------------------- check
def cmd_check(a):
    c = load_course(a.course); n_slides = len(glob.glob(os.path.join(c['dir'], 'slides', 's-*.png')))
    shown, problems, words = set(), [], 0
    for m in modules(c):
        name = os.path.basename(m); bl = parse(m); last = 0
        if not any(k == 'slide' for k, _ in bl): problems.append(f'{name}: no [SLIDE n] marker')
        for k, v in bl:
            if k == 'slide':
                if not 1 <= v <= n_slides: problems.append(f'{name}: [SLIDE {v}] - the deck has {n_slides} slides')
                if v < last: problems.append(f'{name}: [SLIDE {v}] comes after [SLIDE {last}] (going back is allowed, just checking)')
                shown.add(v); last = v
            elif k == 'clip' and not os.path.exists(os.path.join(c['dir'], 'media', v)): problems.append(f'{name}: media/{v} not found')
            elif k == 'say':
                words += len(v.split())
                for sym in re.findall(r'[%&/=<>→€$£#@]', plain(v)):
                    problems.append(f'{name}: "{sym}" in "{v[:60]}..." - spell it out the way it should be spoken'); break
            if k in ('say', 'head') and re.search(r'\{(?![\w-]+\|)|\{[\w-]+\|[^}]*$', v if k == 'say' else v[1]):
                problems.append(f'{name}: unbalanced {{lang|...}} in "{(v if k == "say" else v[1])[:60]}"')
    hidden = {s['slide'] for s in json.load(open(os.path.join(c['dir'], 'source', 'slides.json'), encoding='utf-8')) if s.get('hidden')}
    missing = [i for i in range(1, n_slides + 1) if i not in shown and i not in hidden]
    if missing: problems.append(f'slides never shown: {missing}')
    print('\n'.join(problems) if problems else 'all good')
    print(f'{len(modules(c))} modules, {words} words, about {words / 150:.0f} minutes of speech')

# ---------------------------------------------------------------- speech engines
class Speaker:
    def __init__(self, c):
        self.c, self.engine, self.voice = c, c['engine'], c['voice']
        self.speed, self.lang = float(c.get('speed', 1.0)), c.get('language', 'en')
        self.tmp = tempfile.mkdtemp(); self.vb = {}
        if self.engine == 'kokoro':
            from kokoro_onnx import Kokoro
            base = 'https://github.com/thewh1teagle/kokoro-onnx/releases/download/model-files-v1.0/'
            files = [os.path.join(MODELS, 'kokoro', f) for f in ('kokoro-v1.0.onnx', 'voices-v1.0.bin')]
            for f in files:
                if not os.path.exists(f): download(base + os.path.basename(f), f)
            self.tts = Kokoro(*files)
        elif self.engine == 'piper':
            from piper import PiperVoice
            onnx = os.path.join(MODELS, 'piper', self.voice + '.onnx')
            if not os.path.exists(onnx):
                os.makedirs(os.path.dirname(onnx), exist_ok=True)
                run(sys.executable, '-m', 'piper.download_voices', self.voice, '--data-dir', os.path.dirname(onnx))
            self.tts = PiperVoice.load(onnx)
        elif self.engine != 'voicebox': sys.exit(f'unknown engine {self.engine} (kokoro, piper or voicebox)')
        self.tag = f"{self.engine}|{self.voice}|{self.speed}|{c.get('voicebox_profile')}|{FOREIGN_GAP}|"

    def segments(self, text):
        out, pos = [], 0
        for m in re.finditer(r'\{([\w-]+)\|([^{}]*)\}', text):
            if m.start() > pos: out.append((self.lang, text[pos:m.start()]))
            out.append((m[1], m[2])); pos = m.end()
        if pos < len(text): out.append((self.lang, text[pos:]))
        return [(l, t.strip()) for l, t in out if t.strip()]

    def one(self, text, lang):
        f = os.path.join(self.tmp, 's.wav')
        if self.engine == 'kokoro':
            import soundfile as sf
            code = KOKORO_LANG.get(lang, 'en-gb')
            if lang == 'en' and self.voice[:1] == 'a': code = 'en-us'
            audio, sr = self.tts.create(text, voice=self.voice, speed=self.speed, lang=code); sf.write(f, audio, sr)
        elif self.engine == 'piper':                 # one voice; foreign phrases keep its pronunciation
            from piper import SynthesisConfig
            with wave.open(f, 'wb') as w: self.tts.synthesize_wav(text, w, syn_config=SynthesisConfig(length_scale=1 / self.speed))
        else: return self.voicebox(text, lang)
        return decode(f)

    def speak(self, text):
        import numpy as np
        segs, parts = self.segments(text), []
        for i, (l, t) in enumerate(segs):
            if i:
                g = FOREIGN_GAP if (l != self.lang or segs[i - 1][0] != self.lang) else 0.08
                parts.append(np.zeros((int(g * SR), 2), np.float32))
            parts.append(self.one(t, l))
        return np.concatenate(parts) if parts else np.zeros((1, 2), np.float32)

    def voicebox(self, text, lang):
        import urllib.request, urllib.error
        base = self.c.get('voicebox_url', 'http://127.0.0.1:17493').rstrip('/')
        get = lambda p: json.loads(urllib.request.urlopen(base + p, timeout=30).read())
        if not self.vb:
            profs = get('/profiles'); profs = profs.get('profiles', profs) if isinstance(profs, dict) else profs
            want = (self.c.get('voicebox_profile') or self.voice or '').lower()
            hit = next((p for p in profs if want in (str(p.get('id', '')).lower(), str(p.get('name', '')).lower())), None)
            if not hit: sys.exit(f'Voicebox profile "{want}" not found. Profiles: {[p.get("name") for p in profs]}')
            fields = {}
            try:
                api = get('/openapi.json'); ref = api['paths']['/generate']['post']['requestBody']['content']['application/json']['schema']
                while '$ref' in ref: ref = api['components']['schemas'][ref['$ref'].split('/')[-1]]
                fields = ref.get('properties', {})
            except Exception: pass
            pick = lambda names, d: next((n for n in names if n in fields), d)
            self.vb = dict(profile=hit.get('id', hit.get('name')), text=pick(['text', 'input', 'prompt'], 'text'),
                           prof=pick(['profile_id', 'voice_profile_id', 'profile', 'voice_id', 'voice'], 'profile_id'),
                           lang=pick(['language', 'lang', 'language_code', 'locale'], 'language'))
        body = {self.vb['text']: text, self.vb['prof']: self.vb['profile'], self.vb['lang']: lang}
        req = urllib.request.Request(base + '/generate', data=json.dumps(body).encode(), headers={'Content-Type': 'application/json'})
        try:
            with urllib.request.urlopen(req, timeout=900) as r: data, ctype = r.read(), r.headers.get('Content-Type', '')
        except urllib.error.HTTPError as e: sys.exit(f'Voicebox /generate refused: {e.code} {e.read()[:300]!r}')
        if 'json' in ctype:
            j = json.loads(data); j = j.get('generation', j) if isinstance(j, dict) else j
            ref = next((j[k] for k in ('audio_url', 'url', 'audio_path', 'path', 'file', 'audio') if k in j), None)
            if ref is None and 'id' in j: ref = f"/audio/{j['id']}"
            if ref is None: sys.exit(f'Voicebox answered {j}')
            if str(ref).startswith('/') and not os.path.exists(str(ref)): ref = base + ref
            data = urllib.request.urlopen(ref, timeout=300).read() if str(ref).startswith('http') else open(ref, 'rb').read()
        f = os.path.join(self.tmp, 'vb.audio'); open(f, 'wb').write(data); return decode(f)

def cmd_voicetest(a):
    import soundfile as sf
    c = load_course(a.course); sp = Speaker(c)
    text = {'sv': 'Hej och välkommen till kursen. Så här låter rösten när den läser ett vanligt stycke text.',
            }.get(c['language'], 'Hello, and welcome to the course. This is how the narrator sounds when reading an ordinary paragraph.')
    f = out(c, 'voice_test.wav'); sf.write(f, sp.speak(text), SR); print('wrote', f)

# ---------------------------------------------------------------- audio
def build_audio(c, m, sp, a):
    import numpy as np, soundfile as sf
    name = os.path.basename(m)[:-3]; base = out(c, 'audio', name); blocks = parse(m)
    cache = os.path.join(c['dir'], '.cache', 'tts'); os.makedirs(cache, exist_ok=True)
    files = {}
    for i, (k, v) in enumerate(blocks):                       # 1. every paragraph -> cached mono FLAC
        if k not in ('say', 'head'): continue
        text = v if k == 'say' else v[1] + '.'
        p = os.path.join(cache, hashlib.sha1((sp.tag + text).encode()).hexdigest() + '.flac'); files[i] = p
        if os.path.exists(p): continue
        if over_budget(a): raise Paused(f'{name}: paused during speech, paragraph {i + 1} of {len(blocks)}')
        x = sp.speak(text).mean(axis=1); part = f'{p}.{os.getpid()}.part'
        sf.write(part, x, SR, subtype='PCM_16', format='FLAC'); os.replace(part, p)
    tmp = tempfile.mkdtemp(); allf = os.path.join(tmp, 'all.wav')
    with sf.SoundFile(allf, 'w', SR, 1, 'PCM_16') as f:
        for p in files.values(): f.write(sf.read(p, dtype='float32')[0])
    gain = 10 ** ((SPEECH_LUFS - lufs_of(allf)) / 20)
    load = lambda p: np.repeat(sf.read(p, dtype='float32')[0][:, None], 2, axis=1)
    wav = os.path.join(tmp, 'mix.wav'); t, chapters, timing = 0.0, [], []
    with sf.SoundFile(wav, 'w', SR, 2, 'PCM_16') as mix:      # 2. assemble
        def add(x):
            nonlocal t; mix.write(np.clip(x, -1, 1)); t += len(x) / SR
        sil = lambda s: np.zeros((int(s * SR), 2), np.float32)
        for i, (k, v) in enumerate(blocks):
            if k == 'head':
                chapters.append((t, plain(v[1]))); t0 = t; add(load(files[i]) * gain)
                timing.append(dict(kind='head', level=v[0], start=round(t0, 3), end=round(t, 3), text=plain(v[1]))); add(sil(GAP_HEAD))
            elif k == 'say':
                t0 = t; add(load(files[i]) * gain); timing.append(dict(kind='say', start=round(t0, 3), end=round(t, 3), text=plain(v))); add(sil(GAP_PARA))
            elif k == 'pause': add(sil(v))
            elif k == 'slide': timing.append(dict(kind='slide', slide=v, start=round(t, 3)))
            elif k == 'clip':
                f = os.path.join(c['dir'], 'media', v); x = decode(f)
                cl = os.path.join(tmp, 'clip.wav'); sf.write(cl, x, SR); x = x * 10 ** ((SPEECH_LUFS + 2 - lufs_of(cl)) / 20)
                add(sil(GAP_CLIP - GAP_PARA)); t0 = t; add(x); timing.append(dict(kind='clip', start=round(t0, 3), end=round(t, 3), text=v)); add(sil(GAP_CLIP))
    meta = os.path.join(tmp, 'meta.txt')                      # 3. MP3 + M4B with chapters
    with open(meta, 'w', encoding='utf-8') as f:
        f.write(f";FFMETADATA1\ntitle={c.get('title', '')}: {chapters[0][1] if chapters else name}\nartist={c.get('author', '')}\nalbum={c.get('title', '')}\ngenre=Audiobook\n")
        for (st, nm), en in zip(chapters, [x[0] for x in chapters[1:]] + [t]):
            f.write(f'[CHAPTER]\nTIMEBASE=1/1000\nSTART={int(st * 1000)}\nEND={int(en * 1000)}\ntitle={nm}\n')
    run('ffmpeg', '-y', '-loglevel', 'error', '-i', wav, '-codec:a', 'libmp3lame', '-b:a', '160k', base + '.mp3.part.mp3')
    run('ffmpeg', '-y', '-loglevel', 'error', '-i', wav, '-i', meta, '-map_metadata', '1', '-map_chapters', '1', '-c:a', 'aac', '-b:a', '128k', '-f', 'mp4', base + '.m4b.part')
    json.dump(timing, open(base + '.timing.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
    os.replace(base + '.mp3.part.mp3', base + '.mp3'); os.replace(base + '.m4b.part', base + '.m4b')
    shutil.rmtree(tmp, ignore_errors=True)
    print(f'{name}: {t / 60:.1f} min -> out/audio/{name}.m4b / .mp3', flush=True)

def audio_done(c, m):
    b = out(c, 'audio', os.path.basename(m)[:-3]); return all(os.path.exists(b + e) for e in ('.m4b', '.mp3', '.timing.json')) and \
        os.path.getmtime(b + '.timing.json') >= os.path.getmtime(m)

def cmd_audio(a):
    c = load_course(a.course); todo = [m for m in modules(c, a.only) if not audio_done(c, m)]
    if not todo: print('audio: everything is up to date'); return
    sp = Speaker(c)
    for m in todo: build_audio(c, m, sp, a)

# ---------------------------------------------------------------- video
def hms(t):
    t = int(t); return f'{t // 3600}:{t % 3600 // 60:02d}:{t % 60:02d}' if t >= 3600 else f'{t // 60}:{t % 60:02d}'

def build_video(c, m, a):
    name = os.path.basename(m)[:-3]; ab = out(c, 'audio', name); mp4 = out(c, 'video', name + '.mp4')
    timing = json.load(open(ab + '.timing.json', encoding='utf-8'))
    turns = [(e['start'], e['slide']) for e in timing if e['kind'] == 'slide']
    if not turns: sys.exit(f'{name}: no [SLIDE n] in the script')
    turns[0] = (0.0, turns[0][1])
    turns = [x for i, x in enumerate(turns) if i == len(turns) - 1 or turns[i + 1][0] > x[0]]     # same moment: last one wins
    dur = float(subprocess.run(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'csv=p=0', ab + '.m4b'],
                               capture_output=True, text=True).stdout)
    seg_dir = os.path.join(c['dir'], '.cache', 'video', name); os.makedirs(seg_dir, exist_ok=True)
    bounds = [round(t * FPS) for t, _ in turns] + [round(dur * FPS) + FPS]; segs = []
    vf = 'scale=1920:1080:force_original_aspect_ratio=decrease,pad=1920:1080:(ow-iw)/2:(oh-ih)/2:color=white,format=yuv420p'
    for i, (t, s) in enumerate(turns):
        n = bounds[i + 1] - bounds[i]
        if n <= 0: continue
        png = os.path.join(c['dir'], 'slides', f's-{s:03d}.png')
        seg = os.path.join(seg_dir, f'{i:04d}_{s}_{n}_{int(os.path.getmtime(png))}.mp4'); segs.append(seg)
        if os.path.exists(seg): continue
        if over_budget(a): raise Paused(f'{name}: paused at slide change {i + 1} of {len(turns)}')
        run('ffmpeg', '-y', '-loglevel', 'error', '-loop', 1, '-framerate', FPS, '-i', png, '-frames:v', n, '-vf', vf,
            '-c:v', 'libx264', '-preset', 'veryfast', '-tune', 'stillimage', '-crf', 20, '-f', 'mp4', seg + '.part')
        os.replace(seg + '.part', seg)
    lst = os.path.join(seg_dir, 'list.txt'); open(lst, 'w').write(''.join(f"file '{x}'\n" for x in segs))
    run('ffmpeg', '-y', '-loglevel', 'error', '-f', 'concat', '-safe', 0, '-i', lst, '-i', ab + '.m4b', '-map', '0:v', '-map', '1:a',
        '-c', 'copy', '-movflags', '+faststart', '-shortest', '-f', 'mp4', mp4 + '.part')
    os.replace(mp4 + '.part', mp4)
    marks = []                                                # YouTube chapters: from the headings, each at least 10 s
    for e in timing:
        if e['kind'] != 'head': continue
        st = 0.0 if not marks else e['start']
        if marks and st - marks[-1][0] < 10: continue
        marks.append((st, f"{hms(st)} {e['text']}"))
    open(out(c, 'video', name + '.youtube.txt'), 'w', encoding='utf-8').write('\n'.join(x for _, x in marks) + '\n')
    print(f'{name}: {hms(dur)}, {len(turns)} slide changes -> out/video/{name}.mp4', flush=True)

def video_done(c, m):
    name = os.path.basename(m)[:-3]; v = out(c, 'video', name + '.mp4'); t = out(c, 'audio', name + '.timing.json')
    return os.path.exists(v) and os.path.exists(t) and os.path.getmtime(v) >= os.path.getmtime(t) and \
        subprocess.run(['ffprobe', '-v', 'error', v], capture_output=True).returncode == 0

def cmd_video(a):
    c = load_course(a.course)
    for m in modules(c, a.only):
        if not audio_done(c, m): print(f'{os.path.basename(m)}: run "audio" first'); continue
        if not video_done(c, m): build_video(c, m, a)
    print('video: done')

# ---------------------------------------------------------------- subtitles
MAXLINE, MAXCUE, MINDUR = 42, 84, 1.0

def chunks(text):
    sents = re.findall(r'[^.!?…:;]+[.!?…:;]*["»”)]*\s*', text) or [text]; outl = []
    for s in (x.strip() for x in sents if x.strip()):
        if len(s) > MAXCUE:
            n = -(-len(s) // MAXCUE); cuts = [0]
            for k in range(1, n):
                target = len(s) * k // n
                near = [m.end() for m in re.finditer(r'(?<=\w\w\w\w)[,–—]\s', s) if abs(m.end() - target) <= 14 and m.end() > cuts[-1] + 20]
                sp = [m.end() for m in re.finditer(' ', s) if m.end() > cuts[-1] + 10]
                cuts.append(min(near, key=lambda p: abs(p - target)) if near else min(sp, key=lambda p: abs(p - target)) if sp else target)
            outl.extend(p for p in (s[x:y].strip() for x, y in zip(cuts, cuts[1:] + [len(s)])) if p); continue
        if outl and len(outl[-1]) + 1 + len(s) <= MAXCUE * 0.6: outl[-1] += ' ' + s
        else: outl.append(s)
    return outl

def two_lines(c):
    if len(c) <= MAXLINE: return c
    sp = [m.start() for m in re.finditer(' ', c)]
    i = min(sp, key=lambda p: abs(p - len(c) // 2)) if sp else len(c) // 2
    return c[:i].strip() + '\n' + c[i:].strip()

def ts(t):
    ms = int(round(t * 1000)); h, ms = divmod(ms, 3600000); mi, ms = divmod(ms, 60000); s, ms = divmod(ms, 1000)
    return f'{h:02d}:{mi:02d}:{s:02d},{ms:03d}'

def read_tr(c, lang, name, src):
    """Translations: out/subtitles/tr/LANG/NAME.txt (or .part1.txt, .part2.txt ...) in '@@ id' format."""
    d = out(c, 'subtitles', 'tr', lang, 'x')[:-1]
    parts = sorted(glob.glob(d + name + '.part*.txt'), key=lambda p: int(re.search(r'part(\d+)', p)[1])) or \
        ([d + name + '.txt'] if os.path.exists(d + name + '.txt') else [])
    if not parts: return None, 'no translation yet'
    text = '\n'.join(open(f, encoding='utf-8').read() for f in parts); e = {}
    for m in re.finditer(r'^@@\s*(\d+)\s*\n(.*?)(?=^@@\s*\d+\s*$|\Z)', text, flags=re.S | re.M): e[int(m[1])] = ' '.join(m[2].split())
    miss = [b['id'] for b in src if not e.get(b['id'])]
    if miss: return None, f'missing ids {miss[:20]}'
    return [dict(b, text=e[b['id']]) for b in src], 'ok'

def cmd_subs(a):
    c = load_course(a.course); src_lang = c.get('language', 'en')
    for m in modules(c, a.only):
        name = os.path.basename(m)[:-3]
        src = [dict(id=i, kind=k, text=plain(v if k == 'say' else v[1])) for i, (k, v) in
               enumerate([b for b in parse(m) if b[0] in ('say', 'head')])]
        json.dump(src, open(out(c, 'subtitles', 'src', name + '.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
        if not audio_done(c, m): print(f'{name}: text exported; run "audio" for timed subtitles'); continue
        timing = json.load(open(out(c, 'audio', name + '.timing.json'), encoding='utf-8'))
        spoken = [e for e in timing if e['kind'] in ('head', 'say')]
        for lang in c.get('subtitles', [src_lang]):
            blocks, status = (src, 'ok') if lang == src_lang else read_tr(c, lang, name, src)
            if blocks is None: print(f'{name} {lang}: {status}'); continue
            if len(blocks) != len(spoken): print(f'{name} {lang}: script changed since the audio - run "audio" again'); continue
            cues, bi = [], 0
            for e in timing:
                if e['kind'] == 'clip': cues.append((e['start'], min(e['end'], e['start'] + 6), f"♪ {TRACK_WORD.get(lang, 'audio')}")); continue
                if e['kind'] not in ('head', 'say'): continue
                parts = chunks(blocks[bi]['text']); bi += 1; total = sum(len(p) for p in parts) or 1
                t, d = e['start'], e['end'] - e['start']
                for p in parts:
                    dd = d * len(p) / total; cues.append((t, t + max(dd, min(MINDUR, d)), p)); t += dd
            cues = [(x, min(y, cues[i + 1][0]) if i + 1 < len(cues) else y, z) for i, (x, y, z) in enumerate(cues)]
            with open(out(c, 'subtitles', 'srt', f'{name}.{lang}.srt'), 'w', encoding='utf-8') as f:
                for i, (x, y, z) in enumerate(cues, 1): f.write(f'{i}\n{ts(x)} --> {ts(y)}\n{two_lines(z)}\n\n')
            print(f'{name} {lang}: {len(cues)} subtitles')

def cmd_all(a):
    cmd_audio(a); cmd_video(a); cmd_subs(a)


# ---------------------------------------------------------------- setup and skill
def cmd_setup(a):
    """Download voice models once, so later runs (also from Claude's sandbox) need no network.
    en           Kokoro: one model file with all its voices (English, plus a few other languages)
    sv           Piper's Swedish default voice (sv_SE-nst-medium)
    piper:VOICE  any Piper voice, e.g. piper:de_DE-thorsten-medium"""
    import soundfile as sf
    for item in (a.langs or ['en']):
        if item.startswith('piper:'):
            lang, cfg = item, dict(engine='piper', voice=item.split(':', 1)[1], speed=1.0, language='xx')
        else:
            lang, cfg = item, dict(DEFAULTS.get(item, DEFAULTS['en']), language=item)
        print(f'{lang}: {cfg["engine"]} voice {cfg["voice"]} ...', flush=True)
        sp = Speaker(cfg)
        text = {'sv': 'Hej! Rösten är klar att använda.', 'en': 'Hello! The voice is ready to use.'}.get(item, 'Hello! The voice is ready to use.')
        f = os.path.join(MODELS, f'voice_check_{lang.replace(":", "_")}.wav'); sf.write(f, sp.speak(text), SR)
        print(f'{lang}: ready - listen to {f}')

def cmd_skill(a):
    """Package SKILL.md as a skill: a zip for the Claude apps, and optionally a copy for Claude Code."""
    import zipfile
    name = 'slide-video-course'
    src = os.path.join(KIT, 'SKILL.md')
    os.makedirs(os.path.join(KIT, 'dist'), exist_ok=True)
    z = os.path.join(KIT, 'dist', name + '.zip')
    with zipfile.ZipFile(z, 'w', zipfile.ZIP_DEFLATED) as f:
        f.write(src, f'{name}/SKILL.md')
    print(f'wrote {z}\n  Claude apps: Customize > Skills > Add, choose this file, then switch the skill on.')
    if a.code:
        d = os.path.join(os.path.expanduser('~'), '.claude', 'skills', name); os.makedirs(d, exist_ok=True)
        shutil.copyfile(src, os.path.join(d, 'SKILL.md'))
        print(f'installed for Claude Code: {d}/SKILL.md')

# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    p = sub.add_parser('init'); p.add_argument('deck'); p.add_argument('--course'); p.add_argument('--lang', default='en')
    p = sub.add_parser('setup'); p.add_argument('langs', nargs='*')
    p = sub.add_parser('skill'); p.add_argument('--code', action='store_true')
    for nm in ('check', 'voicetest', 'audio', 'video', 'subs', 'all'):
        p = sub.add_parser(nm); p.add_argument('course'); p.add_argument('only', nargs='*'); p.add_argument('--budget', type=float)
    a = ap.parse_args()
    try: globals()['cmd_' + a.cmd](a)
    except Paused as e: print(e, '- run the same command again to continue'); sys.exit(3)

if __name__ == '__main__':
    main()
